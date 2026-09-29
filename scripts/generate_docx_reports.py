"""
Script to generate 13 separate professional Word (.docx) reports for Weeks 3 to 15
of the Jenkins-Based Podcast Episode Release System.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs")

PRIMARY_COLOR = RGBColor(27, 54, 93)      # Deep Navy #1B365D
SECONDARY_COLOR = RGBColor(74, 96, 122)   # Slate Blue #4A607A
TEXT_COLOR = RGBColor(40, 40, 40)         # Charcoal #282828
MUTED_COLOR = RGBColor(100, 100, 100)     # Gray #646464
CODE_BG = "F4F6F9"
HEADER_BG = "1B365D"
ZEBRA_BG = "F8FAFC"
BORDER_COLOR = "D0D7DE"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def style_table(table, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        # Prevent row split across pages
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if i == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            for cell in row.cells:
                set_cell_background(cell, HEADER_BG)
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9.5)
        else:
            bg = ZEBRA_BG if i % 2 == 1 else "FFFFFF"
            for cell in row.cells:
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for run in p.runs:
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = TEXT_COLOR

    # Set column widths if provided
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = Inches(width)

def add_callout(doc, text, title="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "EBF3FB")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    # Left border only (thick blue)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="1B365D"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"[{title}] ")
    run_title.font.bold = True
    run_title.font.size = Pt(9.5)
    run_title.font.color.rgb = PRIMARY_COLOR
    
    run_text = p.add_run(text)
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = TEXT_COLOR
    doc.add_paragraph() # spacing

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, CODE_BG)
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="6" w:space="0" w:color="{BORDER_COLOR}"/><w:top w:val="single" w:sz="6" w:space="0" w:color="{BORDER_COLOR}"/><w:right w:val="single" w:sz="6" w:space="0" w:color="{BORDER_COLOR}"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="{BORDER_COLOR}"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph() # spacing

def build_base_document(week_num, week_title, summary_desc):
    doc = docx.Document()
    
    # Page Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"Jenkins Podcast Release System | Week {week_num} Milestone Report")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = MUTED_COLOR
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("DevOps Engineering Practicum — 15-Week Curriculum")
        frun.font.name = "Arial"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = MUTED_COLOR

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run(f"Week {week_num}: {week_title}")
    run_title.font.name = "Arial"
    run_title.font.bold = True
    run_title.font.size = Pt(20)
    run_title.font.color.rgb = PRIMARY_COLOR
    
    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run(f"Technical Milestone Report & Verification Record — {summary_desc}")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = SECONDARY_COLOR

    # Metadata Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_data = [
        ("Project Title", "Jenkins-Based Podcast Episode Release System"),
        ("Milestone", f"Week {week_num} — {week_title}"),
        ("Repository", "https://github.com/Talha905/podcast-episode-release-system"),
        ("Branch / Baseline", "origin/main"),
        ("Verification Status", "Fully Implemented, Executed & Verified")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        
        p0 = cell_lbl.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = PRIMARY_COLOR
        
        p1 = cell_val.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = TEXT_COLOR
        
    style_table(meta_table, col_widths=[2.2, 4.3])
    doc.add_paragraph()
    return doc

def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = PRIMARY_COLOR
    return p

def add_body_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = "Arial"
        r_bold.font.bold = True
        r_bold.font.size = Pt(10)
        r_bold.font.color.rgb = TEXT_COLOR
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = TEXT_COLOR
    return p

def add_bullet_point(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = "Arial"
        r_bold.font.bold = True
        r_bold.font.size = Pt(10)
        r_bold.font.color.rgb = TEXT_COLOR
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = TEXT_COLOR
    return p

print("Helper functions defined.")
